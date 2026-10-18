#!/usr/bin/env python3
"""우체국보험 상품 엑셀 → src/config/products.data.json 변환.

사용: python3 scripts/import-products.py <상품 엑셀 경로>
- '본상품' 시트에서 판매상태가 '판매'인 행만 반영합니다. (판매중지·'온라인' 상품 제외)
- 엑셀에 적힌 값만 옮깁니다. 비어 있는 칸(가입나이/보험기간·납입기간/가입한도/주요보장)은 JSON 에 넣지 않으며,
  화면에서는 값이 있을 때만 보입니다. 엑셀에 값을 채운 뒤 이 스크립트를 다시 실행하세요.
"""
import json, sys, pathlib
import openpyxl

COLS = {"id": "ID", "group": "분류", "name": "상품명", "effectiveDate": "고시 시행일", "status": "판매상태",
        "ageRange": "가입나이", "period": "보험기간/납입기간", "limit": "가입한도", "coverage": "주요보장", "job": "직업·직종 확인"}

def main(src: str) -> None:
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["본상품"]
    head = [c.value for c in ws[1]]
    idx = {k: head.index(v) for k, v in COLS.items()}
    out, skipped = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        get = lambda k: (str(row[idx[k]]).strip() if row[idx[k]] not in (None, "") else "")
        if not get("id"):
            continue
        if get("status") != "판매" or "온라인" in get("name"):
            skipped.append((get("id"), get("name"), get("status")))
            continue
        item = {"id": get("id"), "name": get("name"), "group": get("group"), "effectiveDate": get("effectiveDate")}
        for k in ("ageRange", "period", "limit"):
            if get(k): item[k] = get(k)
        if get("coverage"):
            item["coverage"] = [s.strip() for s in get("coverage").replace("\r", "").split("\n") if s.strip()]
        if "직업급수" in get("job"):
            item["jobCheck"] = True
        out.append(item)
    dest = pathlib.Path(__file__).resolve().parent.parent / "src" / "config" / "products.data.json"
    dest.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"반영 {len(out)}개 → {dest}")
    print(f"제외 {len(skipped)}개:", ", ".join(f"{i} {n}({s})" for i, n, s in skipped))

if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
