"""명함 장면 모델 — 도형·글을 한 번만 정의하면 SVG 와 PPTX 가 같은 모양으로 나옵니다.
좌표 단위는 0.1mm (명함 한 면 = 900 x 500 = 90 x 50mm). 글자 크기는 pt.
요소: rect / ellipse / line / poly / text  (모두 dict, 색은 '#RRGGBB')"""
import html
import math

W, H = 900, 500
PT = 3.5278  # 1pt = 0.3528mm = 3.5278 단위
SANS_CSS = "Pretendard,'Noto Sans KR','Noto Sans CJK KR','Malgun Gothic','Apple SD Gothic Neo',sans-serif"
SERIF_CSS = "'Noto Serif KR','Noto Serif CJK KR','Batang','Nanum Myeongjo',serif"
FONT_PPT = {"sans": "맑은 고딕", "serif": "바탕"}
LINE = 1.45  # 줄 간격(글자 크기 배수)


def rect(x, y, w, h, fill=None, stroke=None, sw=0, r=0, op=1, rot=0, dash=False):
    return dict(k="rect", x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, sw=sw, r=r, op=op, rot=rot, dash=dash)


def ellipse(cx, cy, rx, ry, fill=None, stroke=None, sw=0, op=1, rot=0):
    return dict(k="ellipse", cx=cx, cy=cy, rx=rx, ry=ry, fill=fill, stroke=stroke, sw=sw, op=op, rot=rot)


def line(x1, y1, x2, y2, color, sw=2, op=1, dash=False):
    return dict(k="line", x1=x1, y1=y1, x2=x2, y2=y2, color=color, sw=sw, op=op, dash=dash)


def poly(pts, fill=None, stroke=None, sw=0, op=1):
    return dict(k="poly", pts=[(float(a), float(b)) for a, b in pts], fill=fill, stroke=stroke, sw=sw, op=op)


def text(x, y, w, h, s, size, color, bold=False, align="l", font="sans", spacing=0, anchor="t", op=1):
    """상자 (x,y,w,h) 안에 글. align l/c/r, anchor t/m/b, spacing = 자간(pt), 줄바꿈은 \\n"""
    return dict(k="text", x=x, y=y, w=w, h=h, s=s, size=size, color=color, bold=bold, align=align, font=font, spacing=spacing, anchor=anchor, op=op)


def move(els, dx=0, dy=0):
    out = []
    for e in els:
        e = dict(e)
        for a, b in (("x", dx), ("cx", dx), ("x1", dx), ("x2", dx)):
            if a in e:
                e[a] += b
        for a, b in (("y", dy), ("cy", dy), ("y1", dy), ("y2", dy)):
            if a in e:
                e[a] += b
        if "pts" in e:
            e["pts"] = [(p + dx, q + dy) for p, q in e["pts"]]
        out.append(e)
    return out


def scale_about(els, s, ox=0, oy=0):
    """(ox,oy)를 기준점으로 크기 s 배 — 아이콘 배치용(글자 크기는 그대로)."""
    out = []
    for e in els:
        e = dict(e)
        for a in ("x", "cx", "x1", "x2", "w", "rx"):
            if a in e:
                e[a] = e[a] * s
        for a in ("y", "cy", "y1", "y2", "h", "ry"):
            if a in e:
                e[a] = e[a] * s
        if e["k"] == "rect":
            e["r"] = e["r"] * s
        if "sw" in e:
            e["sw"] = e["sw"] * s
        if "pts" in e:
            e["pts"] = [(p * s, q * s) for p, q in e["pts"]]
        out.append(e)
    return move(out, ox, oy)


# ───────────────────────── SVG
def _f(v):
    return f"{v:.1f}".rstrip("0").rstrip(".")


def _paint(e):
    a = []
    a.append(f'fill="{e["fill"]}"' if e.get("fill") else 'fill="none"')
    if e.get("stroke") and e.get("sw"):
        a.append(f'stroke="{e["stroke"]}" stroke-width="{_f(e["sw"])}" stroke-linejoin="round"')
    if e.get("op", 1) != 1:
        a.append(f'opacity="{_f(e["op"])}"')
    return " ".join(a)


def svg_el(e):
    k = e["k"]
    if k == "rect":
        rot = f' transform="rotate({_f(e["rot"])} {_f(e["x"] + e["w"] / 2)} {_f(e["y"] + e["h"] / 2)})"' if e["rot"] else ""
        rr = f' rx="{_f(e["r"])}"' if e["r"] else ""
        dash = ' stroke-dasharray="8 6"' if e.get("dash") else ""
        return f'<rect x="{_f(e["x"])}" y="{_f(e["y"])}" width="{_f(e["w"])}" height="{_f(e["h"])}"{rr} {_paint(e)}{dash}{rot}/>'
    if k == "ellipse":
        rot = f' transform="rotate({_f(e["rot"])} {_f(e["cx"])} {_f(e["cy"])})"' if e["rot"] else ""
        return f'<ellipse cx="{_f(e["cx"])}" cy="{_f(e["cy"])}" rx="{_f(e["rx"])}" ry="{_f(e["ry"])}" {_paint(e)}{rot}/>'
    if k == "line":
        dash = ' stroke-dasharray="8 6"' if e.get("dash") else ""
        op = f' opacity="{_f(e["op"])}"' if e["op"] != 1 else ""
        return f'<line x1="{_f(e["x1"])}" y1="{_f(e["y1"])}" x2="{_f(e["x2"])}" y2="{_f(e["y2"])}" stroke="{e["color"]}" stroke-width="{_f(e["sw"])}" stroke-linecap="round"{dash}{op}/>'
    if k == "poly":
        pts = " ".join(f"{_f(a)},{_f(b)}" for a, b in e["pts"])
        return f'<polygon points="{pts}" {_paint(e)}/>'
    if k == "text":
        fs = e["size"] * PT
        lines = e["s"].split("\n")
        pitch = fs * LINE
        block = pitch * (len(lines) - 1) + fs
        if e["anchor"] == "t":
            y0 = e["y"] + fs * 0.88
        elif e["anchor"] == "m":
            y0 = e["y"] + (e["h"] - block) / 2 + fs * 0.88
        else:
            y0 = e["y"] + e["h"] - block + fs * 0.88
        x = {"l": e["x"], "c": e["x"] + e["w"] / 2, "r": e["x"] + e["w"]}[e["align"]]
        anc = {"l": "start", "c": "middle", "r": "end"}[e["align"]]
        fam = SANS_CSS if e["font"] == "sans" else SERIF_CSS
        sp = f' letter-spacing="{_f(e["spacing"] * PT)}"' if e["spacing"] else ""
        op = f' opacity="{_f(e["op"])}"' if e["op"] != 1 else ""
        ts = "".join(f'<tspan x="{_f(x)}" y="{_f(y0 + i * pitch)}">{html.escape(t)}</tspan>' for i, t in enumerate(lines))
        return f'<text font-family="{fam}" font-size="{_f(fs)}" font-weight="{"700" if e["bold"] else "400"}" fill="{e["color"]}" text-anchor="{anc}"{sp}{op}>{ts}</text>'
    raise ValueError(k)


def to_svg(els, bg, title="명함"):
    body = "".join(svg_el(e) for e in els)
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="90mm" height="50mm" viewBox="0 0 {W} {H}">'
            f'<title>{html.escape(title)}</title><rect width="{W}" height="{H}" fill="{bg}"/>{body}</svg>')


# ───────────────────────── PPTX
def _emu(v):
    return int(round(v * 3600))


def add_to_slide(slide, els, bg):
    from lxml import etree
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE
    from pptx.enum.shapes import MSO_SHAPE as SH
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = RGBColor.from_string(bg.lstrip("#"))

    def clr(h):
        return RGBColor.from_string(h.lstrip("#"))

    def alpha(shape, op):
        if op == 1:
            return
        for sf in shape._element.iter(f"{{{A}}}srgbClr"):
            a = etree.SubElement(sf, f"{{{A}}}alpha"); a.set("val", str(int(op * 100000)))

    def style(shp, e, strokekey="stroke"):
        st = shp._element.find(qn("p:style"))
        if st is not None:
            shp._element.remove(st)
        if e.get("fill"):
            shp.fill.solid(); shp.fill.fore_color.rgb = clr(e["fill"])
        else:
            shp.fill.background()
        if e.get(strokekey) and e.get("sw"):
            shp.line.color.rgb = clr(e[strokekey]); shp.line.width = Emu(_emu(e["sw"]))
            if e.get("dash"):
                shp.line.dash_style = MSO_LINE.DASH
        else:
            shp.line.fill.background()
        alpha(shp, e.get("op", 1))

    for e in els:
        k = e["k"]
        if k == "rect":
            kind = SH.ROUNDED_RECTANGLE if e["r"] else SH.RECTANGLE
            shp = slide.shapes.add_shape(kind, Emu(_emu(e["x"])), Emu(_emu(e["y"])), Emu(_emu(e["w"])), Emu(_emu(e["h"])))
            if e["r"]:
                shp.adjustments[0] = min(0.5, e["r"] / min(e["w"], e["h"]))
            style(shp, e)
            if e["rot"]:
                shp.rotation = e["rot"]
        elif k == "ellipse":
            shp = slide.shapes.add_shape(SH.OVAL, Emu(_emu(e["cx"] - e["rx"])), Emu(_emu(e["cy"] - e["ry"])), Emu(_emu(2 * e["rx"])), Emu(_emu(2 * e["ry"])))
            style(shp, e)
            if e["rot"]:
                shp.rotation = e["rot"]
        elif k == "line":
            from pptx.enum.shapes import MSO_CONNECTOR
            shp = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(_emu(e["x1"])), Emu(_emu(e["y1"])), Emu(_emu(e["x2"])), Emu(_emu(e["y2"])))
            st = shp._element.find(qn("p:style"))
            if st is not None:
                shp._element.remove(st)
            shp.line.color.rgb = clr(e["color"]); shp.line.width = Emu(_emu(e["sw"]))
            if e.get("dash"):
                shp.line.dash_style = MSO_LINE.DASH
            alpha(shp, e["op"])
        elif k == "poly":
            pts = e["pts"]
            fb = slide.shapes.build_freeform(_emu(pts[0][0]), _emu(pts[0][1]), scale=1.0)
            fb.add_line_segments([(_emu(a), _emu(b)) for a, b in pts[1:]], close=True)
            shp = fb.convert_to_shape()
            style(shp, e)
        elif k == "text":
            tb = slide.shapes.add_textbox(Emu(_emu(e["x"])), Emu(_emu(e["y"])), Emu(_emu(e["w"])), Emu(_emu(e["h"])))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[e["anchor"]]
            for i, ln in enumerate(e["s"].split("\n")):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[e["align"]]
                p.line_spacing = Pt(e["size"] * LINE)
                r = p.add_run(); r.text = ln
                r.font.size = Pt(e["size"]); r.font.bold = e["bold"]; r.font.color.rgb = clr(e["color"])
                r.font.name = FONT_PPT[e["font"]]
                rpr = r._r.get_or_add_rPr()
                ea = etree.SubElement(rpr, f"{{{A}}}ea"); ea.set("typeface", FONT_PPT[e["font"]])
                if e["spacing"]:
                    rpr.set("spc", str(int(e["spacing"] * 100)))
                if e["op"] != 1:
                    for sf in rpr.iter(f"{{{A}}}srgbClr"):
                        a = etree.SubElement(sf, f"{{{A}}}alpha"); a.set("val", str(int(e["op"] * 100000)))
            tb.name = "텍스트"


def new_pptx():
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(_emu(W)), Emu(_emu(H))
    return prs


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])
