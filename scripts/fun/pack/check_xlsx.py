"""엑셀 파일 점검: LibreOffice 로 다시 계산한 뒤 수식 오류(#REF! 등)가 없는지 확인합니다.
사용: python3 check_xlsx.py <폴더 또는 파일...> [--show 시트!셀,...]"""
import os, subprocess, sys, tempfile, glob
import warnings
import openpyxl
warnings.filterwarnings("ignore")

ERR = ("#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!", "Err:")


PROFILE = os.path.join(tempfile.gettempdir(), "lo_recalc_profile")


def ensure_profile():
    """LibreOffice 가 불러올 때 항상 다시 계산하도록 설정한 프로필을 만듭니다."""
    d = os.path.join(PROFILE, "user")
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, "registrymodifications.xcu"), "w", encoding="utf-8") as f:
        f.write('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>''')
    return "-env:UserInstallation=file://" + PROFILE


def collect(args):
    files = []
    for a in args:
        files += sorted(glob.glob(os.path.join(a, "**", "*.xlsx"), recursive=True)) if os.path.isdir(a) else [a]
    return files


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    show = []
    if "--show" in sys.argv:
        show = sys.argv[sys.argv.index("--show") + 1].split(","); args = [a for a in args if a != sys.argv[sys.argv.index("--show") + 1]]
    files = collect(args)
    tmp = tempfile.mkdtemp()
    prof = ensure_profile()
    bad = 0
    for k in range(0, len(files), 25):
        chunk = files[k:k + 25]
        subprocess.run(["soffice", prof, "--headless", "--convert-to", "xlsx", "--outdir", tmp] + chunk, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    for f in files:
        out = os.path.join(tmp, os.path.basename(f))
        if not os.path.exists(out):
            print("변환 실패", f); bad += 1; continue
        wb = openpyxl.load_workbook(out, data_only=True)
        errs = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith(ERR):
                        errs.append(f"{ws.title}!{c.coordinate}={c.value}")
        if errs:
            bad += 1; print("오류", os.path.basename(f), errs[:6])
        for s in show:
            sh, cell = s.split("!")
            if sh in wb.sheetnames:
                print(os.path.basename(f), s, wb[sh][cell].value)
    print(f"점검 {len(files)}개, 오류 있는 파일 {bad}개")
    return bad


if __name__ == "__main__":
    sys.exit(1 if main() else 0)
