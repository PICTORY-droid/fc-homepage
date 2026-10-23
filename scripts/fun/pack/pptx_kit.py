"""PPT(python-pptx) 공통 도구 — 테마(색·글꼴) 심기, 10가지 프레임, 감성 모티프, 글·도형·표 도우미.

- 슬라이드는 '제목만' 레이아웃 위에 도형을 얹어 만들어 제목이 개요 보기에 나타납니다.
- 색과 글꼴은 테마에 넣어 두어 PowerPoint 의 [디자인 > 글꼴/색]으로도 바꿀 수 있습니다.
- 사진 자리는 점선 도형이라 도형 채우기 > 그림 으로 사진을 넣을 수 있습니다.
"""
import math
import random
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE as SH
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "맑은 고딕"
AUTHOR = "반딧불이의 곳간"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"

# 테마 번호(0~49) → 모티프 종류
MOTIF_OF = ["flower5", "sun", "petals", "dots", "leaves", "drops", "flower5", "cloud", "leaves", "heart",
            "dots", "petals", "leaves", "drops", "dots", "dots", "sun", "wave", "stars", "cloud",
            "leaves", "wave", "dots", "lantern", "flower5", "flower5", "moon", "dots", "leaves", "leaves",
            "leaves", "stars", "dots", "stars", "dots", "snow", "lantern", "dots", "snow", "dots",
            "heart", "flower5", "lantern", "moon", "lantern", "cloud", "stars", "snow", "moon", "sun"]


def _borders(cell, color, w_pt=0.75):
    """표 칸 테두리(왼·오·위·아래) — 칸 속성에서 채우기보다 앞에 와야 해서 맨 앞에 넣는다."""
    tcPr = cell._tc.get_or_add_tcPr()
    for k, tag in enumerate(("lnL", "lnR", "lnT", "lnB")):
        ln = etree.Element(f"{{{A}}}{tag}", w=str(int(w_pt * 12700)), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, f"{{{A}}}solidFill")
        etree.SubElement(sf, f"{{{A}}}srgbClr").set("val", _hex(color))
        etree.SubElement(ln, f"{{{A}}}prstDash").set("val", "solid")
        tcPr.insert(k, ln)


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def _hex(h):
    return h.lstrip("#").upper()


def _mix(a, b, t):
    a, b = _hex(a), _hex(b)
    return "#%02X%02X%02X" % tuple(round(int(a[i:i + 2], 16) * (1 - t) + int(b[i:i + 2], 16) * t) for i in (0, 2, 4))


class Deck:
    """한 개의 PPT 파일. style: 프레임 번호(0~9), theme_i: 테마 번호(0~49), size: 'wide'(16:9) | 'square'(1:1)"""

    def __init__(self, pal, style, theme_i, size="wide", title="PPT", cover_img=None, modern=False, hero=None, label=""):
        self.pal, self.style, self.ti = pal, style, theme_i
        self.cover_img = cover_img
        self.modern, self.hero, self.label = modern, hero, label
        self.occ = 0  # 같은 구도가 몇 번째로 나오는지(0~4) — 좌우 반전·바탕색을 바꿔 서로 다르게
        self.item = 0
        self.photo_cover = False  # True 면 표지 큰 도형이 사진 자리(그림 채우기로 사진을 넣는 칸)가 된다
        self.card_style, self.badge_style, self.tal = 0, 0, "l"
        self.on_main = "#FFFFFF"  # 강조색 위에 올리는 글자색(새 디자인은 진한 먹색)
        if modern:
            self.on_main = pal["ink"]  # modern: 파스텔·미니멀 새 디자인
        self.prs = Presentation()
        self.W, self.H = (13.333, 7.5) if size == "wide" else (7.5, 7.5)
        self._resize()
        self._patch_theme(title)
        self.rect = (0.7, 1.7, self.W - 0.7, self.H - 0.5)
        self.tbox = (0.7, 0.45, self.W - 2.6, 0.95)
        self.sbox = (0.7, 2.0, 6.0, 1.0)
        self.title_color = pal["dark"]
        self.cal = "l" if cover_img else ("c" if style in (3, 6, 8) else "l")  # 표지 글 정렬
        if modern:
            self.cal = "c" if size != "wide" else "l"
        self.cover_size = (40 if modern else 36) if size == "wide" else 30
        self.prs.core_properties.author = AUTHOR
        self.prs.core_properties.title = title
        self.prs.core_properties.comments = "자유롭게 편집해서 사용하세요."

    # ── 준비
    def _resize(self):
        old_w, old_h = self.prs.slide_width, self.prs.slide_height
        self.prs.slide_width, self.prs.slide_height = Inches(self.W), Inches(self.H)
        fx, fy = self.prs.slide_width / old_w, self.prs.slide_height / old_h
        m = self.prs.slide_master
        for coll in [m.shapes] + [l.shapes for l in m.slide_layouts]:
            for shp in coll:
                if shp.left is not None:
                    shp.left, shp.width = int(shp.left * fx), int(shp.width * fx)
                    shp.top, shp.height = int(shp.top * fy), int(shp.height * fy)

    def _patch_theme(self, name):
        part = self.prs.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        p = self.pal
        root.set("name", name)
        clr = {"dk2": p["dark"], "lt2": p["pale"], "accent1": p["main"], "accent2": p["accent"], "accent3": p["dark"],
               "accent4": p["soft"], "accent5": _mix(p["main"], p["accent"], 0.5), "accent6": _mix(p["dark"], "#FFFFFF", 0.4),
               "hlink": p["main"], "folHlink": p["dark"]}
        for tag, val in clr.items():
            el = root.find(f".//{{{A}}}{tag}")
            if el is not None:
                for ch in list(el):
                    el.remove(ch)
                etree.SubElement(el, f"{{{A}}}srgbClr").set("val", _hex(val))
        for tag in ("majorFont", "minorFont"):
            f = root.find(f".//{{{A}}}{tag}")
            if f is not None:
                f.find(f"{{{A}}}latin").set("typeface", FONT)
                ea = f.find(f"{{{A}}}ea")
                ea.set("typeface", FONT)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # ── 슬라이드
    def slide(self, cover=False):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        self.s = s
        self._cover_now = cover
        self.tal = "l"
        if self.modern:
            self._frame_modern(s, cover)
        elif cover and self.cover_img:
            self._image_cover(s)
        else:
            self._frame(s, cover)
        return s

    def _image_cover(self, s):
        """그림 한 장(우표·폴라로이드·항공우편·승차권)을 깔고 그 위에 제목 글을 올립니다. 글자는 그대로 고칠 수 있어요."""
        W, H = self.W, self.H
        s.shapes.add_picture(self.cover_img, 0, 0, Inches(W), Inches(H)).name = "표지 그림"
        self.title_color = self.pal["dark"]
        if W > H:
            self.tbox = (W * 0.07, H * 0.3, W * 0.5, 1.7)
            self.sbox = (W * 0.07, H * 0.3 + 1.85, W * 0.5, 1.2)
        else:
            self.tbox = (W * 0.1, H * 0.14, W * 0.8, 1.6)
            self.sbox = (W * 0.1, H * 0.14 + 1.75, W * 0.8, 1.2)

    def save(self, path):
        self.prs.save(path)

    # ── 도형·글 도우미 (현재 슬라이드 self.s 에 그린다)
    def shape(self, kind, x, y, w, h, fill=None, line=None, lw=1.0, rot=0, dash=False, name=None, adj=None, s=None):
        s = s or self.s
        shp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        st = shp._element.find(qn("p:style"))
        if st is not None:
            shp._element.remove(st)
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = rgb(line); shp.line.width = Pt(lw)
            if dash:
                shp.line.dash_style = MSO_LINE.DASH
        if rot:
            shp.rotation = rot
        if adj is not None:
            shp.adjustments[0] = adj
        if name:
            shp.name = name
        return shp

    def text(self, x, y, w, h, content, size=16, bold=False, color=None, align="l", anchor="t", italic=False, spacing=1.15, bullet=False, name=None, space_after=0, s=None):
        s = s or self.s
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        self._fill_text(tb.text_frame, content, size, bold, color or self.pal["ink"], align, anchor, italic, spacing, bullet, space_after)
        return tb

    def _fill_text(self, tf, content, size, bold, color, align, anchor, italic, spacing, bullet, space_after=0):
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        lines = content if isinstance(content, (list, tuple)) else [content]
        for k, line in enumerate(lines):
            para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            para.line_spacing = spacing
            if space_after:
                para.space_after = Pt(space_after)
            run = para.add_run()
            run.text = line
            f = run.font
            f.size, f.bold, f.italic = Pt(size), bold, italic
            f.color.rgb = rgb(color)
            if bullet:
                pPr = para._p.get_or_add_pPr()
                pPr.set("marL", "285750"); pPr.set("indent", "-285750")
                bu = etree.SubElement(pPr, f"{{{A}}}buFont"); bu.set("typeface", "Arial")
                ch = etree.SubElement(pPr, f"{{{A}}}buChar"); ch.set("char", "•")

    def title(self, content, size=30, color=None, align="l", box=None):
        t = self.s.shapes.title
        el = t._element
        tree = el.getparent()
        tree.remove(el); tree.append(el)  # 장식 도형 위로 올린다(카드·리본 안에서도 제목이 보이게)
        x, y, w, h = box or self.tbox
        t.left, t.top, t.width, t.height = Inches(x), Inches(y), Inches(w), Inches(h)
        if self.modern and self.tal == "c" and align == "l" and self.s is not None and not getattr(self, "_cover_now", False):
            align = "c"
        self._fill_text(t.text_frame, content, size, True, color or self.title_color, align, "m", False, 1.05, False)
        return t

    def box(self, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.08, dash=False, name=None):
        if self.modern:
            p, cs = self.pal, self.card_style
            r_ = min(0.07, radius)
            if cs == 0:    # 납작한 파스텔 면
                return self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, fill or p["pale"], None, 0, name=name, adj=r_)
            if cs == 1:    # 흰 면 + 얇은 윤곽
                return self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, "#FFFFFF", p["tint"], 1.0, name=name, adj=r_)
            if cs == 2:    # 파스텔 면 + 왼쪽 포인트 막대
                b_ = self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, fill or p["pale"], None, 0, name=name, adj=min(0.04, r_))
                self.shape(SH.RECTANGLE, x, y + 0.12, 0.06, h - 0.24, fill=p["main"], name="포인트 막대")
                return b_
            if cs == 3:    # 윤곽만
                return self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, None, p["main"], 1.25, name=name, adj=r_)
            b_ = self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, "#FFFFFF", p["soft"], 0.75, name=name, adj=min(0.03, r_))  # 흰 면 + 위쪽 색 띠
            self.shape(SH.RECTANGLE, x + 0.01, y + 0.01, w - 0.02, 0.09, fill=p["main"], name="위쪽 띠")
            return b_
        return self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, fill, line, lw, dash=dash, name=name, adj=radius)

    def photo(self, x, y, w, h, label="사진을 넣어 주세요", size=12, shape="round"):
        """사진 자리: 점선 도형. 도형 채우기 > 그림 으로 사진을 넣을 수 있어요."""
        p = self.pal
        kind = SH.OVAL if shape == "oval" else SH.ROUNDED_RECTANGLE
        shp = self.shape(kind, x, y, w, h, fill=p["soft"], line=p["main"], lw=1.5, dash=True, name="사진 자리", adj=0.06 if shape != "oval" else None)
        if self.modern:  # 사진 자리: 파스텔 면 + 작은 글자(점선 없음)
            shp.line.fill.background()
            self._fill_text(shp.text_frame, label, size, False, "#8E8780", "c", "m", False, 1.0, False)
            return shp
        self._fill_text(shp.text_frame, label, size, False, p["dark"], "c", "m", False, 1.0, False)
        return shp

    def table(self, x, y, w, rows, col_w=None, row_h=0.45, size=12, head=True, first_col=False, name="표"):
        p = self.pal
        n, m = len(rows), len(rows[0])
        gf = self.s.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(row_h * n))
        gf.name = name
        tbl = gf.table
        tbl.first_row = head
        tbl.horz_banding = False
        cw = col_w or [w / m] * m
        for j, v in enumerate(cw):
            tbl.columns[j].width = Inches(v)
        for i, row in enumerate(rows):
            tbl.rows[i].height = Inches(row_h)
            for j, v in enumerate(row):
                c = tbl.cell(i, j)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                c.margin_left = c.margin_right = Inches(0.08)
                hd = head and i == 0
                fc = hd or (first_col and j == 0)
                c.fill.solid()
                c.fill.fore_color.rgb = rgb(p["main"] if hd else p["soft"] if fc else ("#FFFFFF" if i % 2 else p["pale"]))
                tf = c.text_frame
                tf.word_wrap = True
                para = tf.paragraphs[0]
                para.alignment = PP_ALIGN.CENTER if (hd or j > 0) else PP_ALIGN.LEFT
                _borders(c, p["soft"] if not hd else p["main"])
                run = para.add_run(); run.text = str(v)
                run.font.size = Pt(size); run.font.bold = bool(hd or fc)
                run.font.color.rgb = rgb(self.on_main if hd else p["dark"] if fc else p["ink"])
        return gf

    # ── 모티프
    def motif(self, cx, cy, size, cols=None, seed=None):
        p = self.pal
        c1, c2, c3 = cols or (p["main"], p["soft"], p["accent"])
        kind = MOTIF_OF[self.ti]
        r = random.Random(self.ti * 31 + (seed or 0))
        add = lambda k, x, y, w, h, f, rot=0: self.shape(k, x, y, w, h, fill=f, rot=rot, name="장식")
        if kind == "flower5":
            for (fx, fy, fs, ps) in ((0, 0, 1.0, 0), (-0.55, 0.45, 0.5, 20), (0.55, 0.5, 0.42, 40)):
                R, pr = size * 0.26 * fs, size * 0.3 * fs
                for k in range(5):
                    a = math.radians(k * 72 + ps)
                    px, py = cx + fx * size + R * math.sin(a), cy + fy * size - R * math.cos(a)
                    add(SH.OVAL, px - pr * 0.45, py - pr * 0.7, pr * 0.9, pr * 1.4, c1 if k % 2 == 0 else c2, k * 72 + ps)
                add(SH.OVAL, cx + fx * size - pr * 0.3, cy + fy * size - pr * 0.3, pr * 0.6, pr * 0.6, c3)
        elif kind == "petals":
            for k in range(8):
                s_ = size * r.uniform(0.18, 0.32)
                add(SH.TEAR, cx + r.uniform(-0.5, 0.5) * size - s_ / 2, cy + r.uniform(-0.5, 0.5) * size - s_ / 2, s_, s_, [c1, c2, c3][k % 3], r.uniform(0, 360))
        elif kind == "sun":
            add(SH.SUN, cx - size * 0.45, cy - size * 0.45, size * 0.9, size * 0.9, c1)
            for k in range(3):
                add(SH.OVAL, cx + r.uniform(-0.6, 0.6) * size, cy + r.uniform(0.45, 0.7) * size, size * 0.1, size * 0.1, c2)
        elif kind == "drops":
            for k in range(7):
                s_ = size * r.uniform(0.14, 0.26)
                add(SH.TEAR, cx + r.uniform(-0.55, 0.55) * size - s_ / 2, cy + r.uniform(-0.55, 0.55) * size - s_ / 2, s_, s_, [c1, c2, c3][k % 3], 315)
        elif kind == "dots":
            for k in range(11):
                s_ = size * r.uniform(0.07, 0.22)
                add(SH.OVAL, cx + r.uniform(-0.55, 0.55) * size - s_ / 2, cy + r.uniform(-0.55, 0.55) * size - s_ / 2, s_, s_, [c1, c2, c3][k % 3])
        elif kind == "leaves":
            for k in range(6):
                lw_, lh = size * 0.17, size * r.uniform(0.42, 0.62)
                add(SH.OVAL, cx + r.uniform(-0.45, 0.45) * size - lw_ / 2, cy + r.uniform(-0.3, 0.3) * size - lh / 2, lw_, lh, [c1, c2, c3][k % 3], r.uniform(-65, 65))
        elif kind == "wave":
            for k, c in enumerate((c2, c1, c3)):
                add(SH.WAVE, cx - size * 0.6, cy - size * 0.3 + k * size * 0.26, size * 1.2, size * 0.3, c)
        elif kind == "stars":
            for k in range(6):
                s_ = size * r.uniform(0.14, 0.34)
                add(SH.STAR_5_POINT, cx + r.uniform(-0.55, 0.55) * size - s_ / 2, cy + r.uniform(-0.55, 0.55) * size - s_ / 2, s_, s_, [c1, c2, c3][k % 3], r.uniform(-20, 20))
        elif kind == "snow":
            for k in range(6):
                s_ = size * r.uniform(0.14, 0.34)
                add(SH.STAR_6_POINT, cx + r.uniform(-0.55, 0.55) * size - s_ / 2, cy + r.uniform(-0.55, 0.55) * size - s_ / 2, s_, s_, [c1, c2, c3][k % 3], r.uniform(0, 60))
        elif kind == "moon":
            add(SH.MOON, cx - size * 0.3, cy - size * 0.5, size * 0.6, size, c1)
            for k in range(3):
                s_ = size * r.uniform(0.1, 0.2)
                add(SH.STAR_5_POINT, cx + r.uniform(0.1, 0.55) * size, cy + r.uniform(-0.5, 0.5) * size, s_, s_, [c2, c3][k % 2])
        elif kind == "cloud":
            add(SH.CLOUD, cx - size * 0.55, cy - size * 0.35, size * 0.95, size * 0.55, c2)
            add(SH.CLOUD, cx - size * 0.1, cy - size * 0.1, size * 0.6, size * 0.38, c1)
            for k in range(3):
                s_ = size * 0.12
                add(SH.TEAR, cx - size * 0.3 + k * size * 0.3, cy + size * 0.32, s_, s_, c3, 315)
        elif kind == "heart":
            for k, (fx, fy, fs, rt) in enumerate(((0, 0, 0.55, 0), (-0.4, 0.3, 0.32, -18), (0.42, 0.28, 0.3, 16))):
                add(SH.HEART, cx + fx * size - size * fs / 2, cy + fy * size - size * fs / 2, size * fs, size * fs, [c1, c2, c3][k], rt)
        elif kind == "lantern":
            for k, fx in enumerate((-0.38, 0.0, 0.38)):
                lw_, lh = size * 0.3, size * (0.5 - 0.08 * (k % 2))
                x0, y0 = cx + fx * size - lw_ / 2, cy - lh / 2 + (k % 2) * size * 0.14
                add(SH.ROUNDED_RECTANGLE, x0, y0, lw_, lh, [c1, c2, c3][k])
                add(SH.RECTANGLE, x0 + lw_ * 0.25, y0 - size * 0.05, lw_ * 0.5, size * 0.05, c3 if k != 2 else c1)
                add(SH.RECTANGLE, x0 + lw_ * 0.25, y0 + lh, lw_ * 0.5, size * 0.05, c3 if k != 2 else c1)

    # ── 새 디자인: 깨끗한 에디토리얼(인스타 감성) — 여백·큰 글자·얇은 선·한 가지 파스텔 면
    def alpha(self, shp, pct):
        """단색 채우기의 투명도(pct = 불투명 %)"""
        cl = shp._element.spPr.find(f"{{{A}}}solidFill/{{{A}}}srgbClr")
        if cl is not None:
            for x in cl.findall(f"{{{A}}}alpha"):
                cl.remove(x)
            etree.SubElement(cl, f"{{{A}}}alpha", val=str(int(pct * 1000)))
        return shp

    def track(self, shp, spc=200):
        """글자 사이 간격(100 = 1pt) — 작은 영문 라벨을 고급스럽게"""
        for para in shp.text_frame.paragraphs:
            for r in para.runs:
                r._r.get_or_add_rPr().set("spc", str(spc))
        return shp

    def _glyph(self, s, kind, cx, cy, d, color):
        """아주 작은 포인트 도형(별·원·더하기·물결·반원) — 한 장에 하나만"""
        k = {"star": SH.STAR_4_POINT, "dot": SH.OVAL, "plus": SH.MATH_PLUS, "ring": SH.DONUT, "heart": SH.HEART, "moon": SH.MOON,
             "sun": SH.SUN, "drop": SH.TEAR, "wave": SH.WAVE, "arch": SH.ROUND_2_SAME_RECTANGLE}[kind]
        shp = self.shape(k, cx - d / 2, cy - d / 2, d, d if kind != "wave" else d * 0.4, fill=color, name="포인트", s=s, rot=315 if kind == "drop" else 0)
        return shp

    GLYPH = ["star", "dot", "plus", "ring", "heart", "moon", "sun", "drop", "wave", "arch"]

    def _frame_modern(self, s, cover):
        p, W, H, v = self.pal, self.W, self.H, self.style % 10
        wide = W > H
        sh = lambda *a, **k: self.shape(*a, s=s, name="장식", **k)
        s.background.fill.solid()
        if cover:
            self._cover_modern(s, v, wide, sh)
            return
        occ = self.occ % 5
        bgc = [p["paper"], "#FFFFFF", p["pale"], p["soft"], "#FFFFFF"][occ]
        s.background.fill.fore_color.rgb = rgb(bgc)
        if occ == 4:
            sh(SH.RECTANGLE, 0, H - 0.32, W, 0.32, fill=p["tint"])
        hd = v % 5
        gk = self.GLYPH[(v + self.occ + 3 * self.item) % 10]
        self.title_color = p["ink"]
        self.rect = (0.7, 1.95, W - 0.7, H - (0.75 if occ == 4 else 0.55))
        if hd == 0:    # 위 얇은 선 + 작은 영문 라벨 + 포인트
            sh(SH.RECTANGLE, 0.7, 0.62, W - 1.4, 0.012, fill=p["ink"])
            self.track(self.text(0.7, 0.26, 6.0, 0.3, self.label, size=9, bold=True, color=p["ink"], anchor="m", name="상단 라벨", s=s), 250)
            self._glyph(s, gk, W - 0.95, 0.4, 0.26, p["main"])
            self.tbox = (0.7, 0.8, W - 1.4, 0.9)
        elif hd == 1:  # 제목 옆 세로 막대
            sh(SH.ROUNDED_RECTANGLE, 0.7, 0.62, 0.07, 0.95, fill=p["main"], adj=0.5)
            self.track(self.text(0.95, 0.3, 6.0, 0.3, self.label, size=9, bold=True, color="#8E8780", anchor="m", name="상단 라벨", s=s), 250)
            self._glyph(s, gk, W - 0.95, 0.9, 0.34, p["tint"])
            self.tbox = (0.95, 0.62, W - 3.0, 0.95)
        elif hd == 2:  # 가운데 정렬 제목 + 작은 포인트 + 짧은 선
            self._glyph(s, gk, W / 2, 0.42, 0.3, p["main"])
            sh(SH.RECTANGLE, W / 2 - 0.45, 1.72, 0.9, 0.025, fill=p["main"])
            self.track(self.text(0.7, 0.2, 4.0, 0.3, self.label, size=8, bold=True, color="#8E8780", anchor="m", name="상단 라벨", s=s), 250)
            self.tbox = (1.5, 0.7, W - 3.0, 0.95)
            self.tal = "c"
        elif hd == 3:  # 아래 얇은 선 + 라벨(아래), 위는 제목만
            sh(SH.RECTANGLE, 0.7, H - 0.5, W - 1.4, 0.012, fill=p["ink"])
            self.track(self.text(0.7, H - 0.46, 6.0, 0.3, self.label, size=9, bold=True, color=p["ink"], anchor="m", name="하단 라벨", s=s), 250)
            self._glyph(s, gk, W - 0.95, H - 0.3, 0.24, p["main"])
            self.tbox = (0.7, 0.55, W - 1.4, 1.0)
            self.rect = (0.7, 1.75, W - 0.7, H - 0.75)
        else:          # 모서리 큰 연한 원 + 라벨
            sh(SH.OVAL, W - 2.2, -1.8, 3.8, 3.8, fill=p["tint"])
            self.track(self.text(0.7, 0.3, 6.0, 0.3, self.label, size=9, bold=True, color=p["ink"], anchor="m", name="상단 라벨", s=s), 250)
            self._glyph(s, gk, W - 1.1, 0.7, 0.3, "#FFFFFF")
            self.tbox = (0.7, 0.72, W - 3.4, 0.95)
        self.card_style = (v // 2) % 5
        self.badge_style = v % 5

    def _photo_spot(self, shp):
        """표지 큰 도형을 사진 자리로 — 도형 서식 > 채우기 > 그림 으로 사진을 넣을 수 있다."""
        shp.name = "사진 자리"
        self._fill_text(shp.text_frame, "사진을 넣어 주세요", 14, False, self.pal["dark"], "c", "m", False, 1.0, False)
        return shp

    def _cover_modern(self, s, v, wide, sh):
        p, W, H, occ = self.pal, self.W, self.H, self.occ
        if self.photo_cover and wide and v not in (0, 2, 4):
            v = (0, 4, 2)[(self.occ + self.ti) % 3]  # 사진 칸이 생기는 구도만 쓴다
        inv = occ in (2, 3) and not self.photo_cover   # 어두운 바탕 표지(사진 표지는 파스텔만)
        mirror = occ in (1, 3)       # 좌우 반전
        tinted = occ == 4            # 파스텔 바탕 표지
        ink = "#FFFFFF" if inv else p["ink"]
        bg = p["dark"] if inv else (p["soft"] if tinted else p["pale"])
        shape_c = p["main"] if not tinted else "#FFFFFF"
        shape2 = p["tint"] if not inv else _mix(p["main"], p["dark"], 0.35)
        s.background.fill.fore_color.rgb = rgb(bg)
        fx = (lambda x, w=0: W - x - w) if mirror else (lambda x, w=0: x)
        lab = f"{self.label}"
        if not wide:  # 정사각(카드뉴스): 위에 큰 도형, 아래 제목
            sh(SH.OVAL, W / 2 - 2.0, 0.9, 4.0, 4.0, fill=shape_c)
            sh(SH.OVAL, W / 2 - 0.2, 0.55, 0.9, 0.9, fill=shape2)
            self.track(self.text(0.8, 5.0, W - 1.6, 0.3, lab, size=10, bold=True, color=ink, align="c", name="라벨", s=s), 250)
            self.tbox, self.sbox = (0.8, 5.35, W - 1.6, 1.0), (0.8, 6.35, W - 1.6, 0.7)
            self.title_color, self.cal = ink, "c"
            self.cover_sub = ink
            return
        tx, ty, tw = 0.9, H * 0.38, W * 0.46
        al = "l"
        if v == 0:   # 아치
            a0 = sh(SH.ROUND_2_SAME_RECTANGLE, fx(W * 0.62, W * 0.26), H * 0.1, W * 0.26, H * 0.8, fill=(p["tint"] if self.photo_cover else shape_c), adj=0.5)
            if self.photo_cover:
                self._photo_spot(a0)
            sh(SH.OVAL, fx(W * 0.56, 1.1), H * 0.62, 1.1, 1.1, fill=shape2)
        elif v == 1:  # 큰 원 두 겹
            sh(SH.OVAL, fx(W * 0.52, 7.2), H * 0.5 - 3.6, 7.2, 7.2, fill=shape2)
            sh(SH.OVAL, fx(W * 0.6, 4.6), H * 0.5 - 2.3, 4.6, 4.6, fill=shape_c)
        elif v == 2:  # 오른쪽 색 면 + 안쪽 얇은 선 프레임
            a0 = sh(SH.RECTANGLE, fx(W * 0.64, W * 0.36), 0, W * 0.36, H, fill=(p["tint"] if self.photo_cover else shape_c))
            if self.photo_cover:
                self._photo_spot(a0)
            sh(SH.RECTANGLE, fx(W * 0.67, W * 0.3), 0.4, W * 0.3, H - 0.8, fill=None, line=ink if not inv else "#FFFFFF", lw=0.75)
        elif v == 3:  # 위쪽 색 면 + 걸친 원
            sh(SH.RECTANGLE, 0, 0, W, H * 0.46, fill=shape_c)
            sh(SH.OVAL, fx(W * 0.72, 2.6), H * 0.46 - 1.3, 2.6, 2.6, fill=shape2)
            tx, ty, tw = 0.9, H * 0.6, W * 0.6
        elif v == 4:  # 어긋난 사각 + 선 프레임
            a0 = sh(SH.ROUNDED_RECTANGLE, fx(W * 0.58, W * 0.3), H * 0.14, W * 0.3, H * 0.66, fill=(p["tint"] if self.photo_cover else shape_c), adj=0.04)
            if self.photo_cover:
                self._photo_spot(a0)
            sh(SH.ROUNDED_RECTANGLE, fx(W * 0.62, W * 0.3), H * 0.18, W * 0.3, H * 0.66, fill=None, line=ink if not inv else "#FFFFFF", lw=0.75, adj=0.04)
        elif v == 5:  # 아치 세 개
            for k, c in enumerate((shape2, shape_c, p["soft"] if not inv else shape2)):
                sh(SH.ROUND_2_SAME_RECTANGLE, fx(W * (0.56 + k * 0.13), W * 0.12), H * (0.3 - 0.07 * k), W * 0.12, H * (0.7 + 0.07 * k), fill=c, adj=0.5)
        elif v == 6:  # 글만 가운데 — 작은 원 하나와 선
            sh(SH.OVAL, W / 2 - 0.45, H * 0.16, 0.9, 0.9, fill=shape_c)
            sh(SH.RECTANGLE, W / 2 - 0.6, H * 0.68, 1.2, 0.02, fill=ink)
            tx, ty, tw, al = W * 0.15, H * 0.34, W * 0.7, "c"
        elif v == 7:  # 왼쪽 색 면
            sh(SH.RECTANGLE, fx(0, W * 0.3), 0, W * 0.3, H, fill=shape_c)
            sh(SH.OVAL, fx(W * 0.3 - 0.6, 1.2), H * 0.5 - 0.6, 1.2, 1.2, fill=shape2)
            tx, ty, tw = W * 0.38, H * 0.38, W * 0.55
            if mirror:
                tx = 0.9
        elif v == 8:  # 모서리 큰 원 + 작은 점
            sh(SH.OVAL, fx(-1.6, 6.0), H - 3.4, 6.0, 6.0, fill=shape_c)
            sh(SH.OVAL, fx(W * 0.82, 0.5), H * 0.2, 0.5, 0.5, fill=shape2)
            sh(SH.OVAL, fx(W * 0.9, 0.2), H * 0.36, 0.2, 0.2, fill=shape_c)
            tx, ty = 0.9 if not mirror else W * 0.42, H * 0.3
        else:         # 반쪽 원
            sh(SH.OVAL, fx(W * 0.58, 7.0), -3.5, 7.0, 7.0, fill=shape2)
            sh(SH.OVAL, fx(W * 0.7, 3.6), -1.0, 3.6, 3.6, fill=shape_c)
        if mirror and v in (0, 1, 2, 4, 5, 9):  # 도형이 왼쪽으로 옮겨 갔으니 글은 오른쪽에
            tx, tw = W * 0.56, W * 0.38
        self.track(self.text(tx, ty - 0.5, tw, 0.3, lab, size=10, bold=True, color=ink, align=al, name="라벨", s=s), 250)
        self.tbox, self.sbox = (tx, ty, tw, 1.3), (tx, ty + 1.4, tw, 1.4)
        self.title_color, self.cal = ink, al
        self.cover_sub = ink

    # ── 프레임(10가지)
    def _frame(self, s, cover):
        p, W, H, st = self.pal, self.W, self.H, self.style
        sh = lambda *a, **k: self.shape(*a, **k)
        wide = W > H
        bg = p["pale"]
        tc = p["dark"]
        tb = (0.7, 0.5, W - 2.7, 0.95)
        rect = (0.7, 1.75, W - 0.7, H - 0.5)
        mo = (W - 0.95, 0.85, 1.1, None)  # 작은 모티프(우상단)
        cv_title = (W * 0.1, H * 0.34, W * 0.8, 1.5)
        cv_sub = (W * 0.1, H * 0.34 + 1.6, W * 0.8, 1.2)
        cv_motif = (W * 0.8, H * 0.22, H * 0.42, None)
        deco = []
        if st == 0:  # 위 띠 / 왼쪽 큰 면
            if cover:
                bg = p["pale"]; deco = [(SH.RECTANGLE, 0, 0, W * 0.38, H, p["main"]), (SH.OVAL, -W * 0.08, H * 0.62, W * 0.3, W * 0.3, _mix(p["main"], "#FFFFFF", 0.18))]
                cv_title = (W * 0.43, H * 0.3, W * 0.52, 1.7); cv_sub = (W * 0.43, H * 0.3 + 1.8, W * 0.52, 1.2)
                cv_motif = (W * 0.19, H * 0.38, W * 0.3, "w")
            else:
                deco = [(SH.RECTANGLE, 0, 0, W, 0.3, p["main"]), (SH.RECTANGLE, 0, 0.3, W, 0.06, p["accent"])]
        elif st == 1:  # 세로 띠
            if cover:
                deco = [(SH.RECTANGLE, 0, 0, W * 0.16, H, p["main"]), (SH.RECTANGLE, W * 0.16, 0, 0.1, H, p["accent"])]
                cv_title = (W * 0.24, H * 0.32, W * 0.62, 1.6); cv_sub = (W * 0.24, H * 0.32 + 1.7, W * 0.62, 1.2)
                cv_motif = (W * 0.8, H * 0.74, H * 0.36, None)
            else:
                deco = [(SH.RECTANGLE, 0, 0, 0.45, H, p["main"]), (SH.RECTANGLE, 0.45, 0, 0.07, H, p["accent"])]
                tb = (1.0, 0.5, W - 3.0, 0.95); rect = (1.0, 1.75, W - 0.7, H - 0.5)
        elif st == 2:  # 큰 원
            bg = "#FFFFFF"
            if cover:
                deco = [(SH.OVAL, W * 0.58, -H * 0.25, H * 1.15, H * 1.15, p["soft"]), (SH.OVAL, W * 0.78, H * 0.58, H * 0.5, H * 0.5, p["main"])]
                cv_title = (W * 0.07, H * 0.3, W * 0.55, 1.6); cv_sub = (W * 0.07, H * 0.3 + 1.7, W * 0.5, 1.2)
                cv_motif = (W * 0.8, H * 0.24, H * 0.42, None)
            else:
                deco = [(SH.OVAL, W - 2.5, -1.6, 4.2, 4.2, p["soft"]), (SH.OVAL, W - 1.2, 0.3, 0.45, 0.45, p["main"])]
        elif st == 3:  # 액자
            if cover:
                deco = [(SH.RECTANGLE, 0.3, 0.3, W - 0.6, H - 0.6, None, p["main"], 3.0), (SH.RECTANGLE, 0.5, 0.5, W - 1.0, H - 1.0, None, p["soft"], 1.0)]
                cv_title = (W * 0.12, H * 0.34, W * 0.76, 1.5); cv_sub = (W * 0.12, H * 0.34 + 1.6, W * 0.76, 1.2)
                cv_motif = (W * 0.5, H * 0.2, H * 0.24, None)
            else:
                deco = [(SH.RECTANGLE, 0.28, 0.28, W - 0.56, H - 0.56, None, p["main"], 2.25), (SH.RECTANGLE, 0.42, 0.42, W - 0.84, H - 0.84, None, p["soft"], 0.75)]
                tb = (0.9, 0.6, W - 3.2, 0.9); rect = (0.9, 1.7, W - 0.9, H - 0.75); mo = (W - 1.25, 1.05, 1.0, None)
        elif st == 4:  # 물결
            bg = "#FFFFFF"
            if cover:
                deco = [(SH.WAVE, -0.3, H * 0.62, W + 0.6, H * 0.4, p["soft"]), (SH.WAVE, -0.3, H * 0.74, W + 0.6, H * 0.34, p["main"])]
                cv_title = (W * 0.08, H * 0.2, W * 0.62, 1.6); cv_sub = (W * 0.08, H * 0.2 + 1.7, W * 0.6, 1.2)
                cv_motif = (W * 0.82, H * 0.28, H * 0.4, None)
            else:
                deco = [(SH.WAVE, -0.2, H - 0.78, W + 0.4, 0.8, p["soft"]), (SH.WAVE, -0.2, H - 0.5, W + 0.4, 0.55, p["main"])]
                rect = (0.7, 1.7, W - 0.7, H - 1.0)
        elif st == 5:  # 대각선(삼각)
            if cover:
                deco = [(SH.RIGHT_TRIANGLE, 0, H * 0.42, W * 0.45, H * 0.58, p["soft"]), (SH.RIGHT_TRIANGLE, 0, H * 0.66, W * 0.26, H * 0.34, p["main"]), (SH.RIGHT_TRIANGLE, W * 0.64, 0, W * 0.36, H * 0.4, p["soft"], None, 1, 180)]
                cv_title = (W * 0.08, H * 0.16, W * 0.62, 1.6); cv_sub = (W * 0.08, H * 0.16 + 1.7, W * 0.6, 1.2)
                cv_motif = (W * 0.84, H * 0.7, H * 0.4, None)
            else:
                deco = [(SH.RIGHT_TRIANGLE, 0, H - 0.9, 1.9, 0.9, p["soft"]), (SH.RIGHT_TRIANGLE, 0, H - 0.5, 1.0, 0.5, p["main"])]
        elif st == 6:  # 카드
            bg = p["dark"] if cover else p["main"]
            if cover:
                deco = [(SH.ROUNDED_RECTANGLE, 0.5, 0.5, W - 1.0, H - 1.0, "#FFFFFF", None, 1, 0, 0.05)]
                cv_title = (W * 0.12, H * 0.33, W * 0.76, 1.5); cv_sub = (W * 0.12, H * 0.33 + 1.6, W * 0.76, 1.2)
                cv_motif = (W * 0.5, H * 0.2, H * 0.22, None)
            else:
                deco = [(SH.ROUNDED_RECTANGLE, 0.3, 0.3, W - 0.6, H - 0.6, "#FFFFFF", None, 1, 0, 0.035)]
                tb = (0.8, 0.6, W - 3.0, 0.95); rect = (0.8, 1.75, W - 0.8, H - 0.7); mo = (W - 1.3, 1.05, 1.0, None)
        elif st == 7:  # 아치
            if cover:
                deco = [(SH.ROUND_2_SAME_RECTANGLE, W * 0.6, H * 0.14, W * 0.3, H * 0.86, p["soft"], None, 1, 0, 0.5), (SH.ROUND_2_SAME_RECTANGLE, W * 0.65, H * 0.26, W * 0.2, H * 0.74, p["main"], None, 1, 0, 0.5)]
                cv_title = (W * 0.07, H * 0.3, W * 0.5, 1.7); cv_sub = (W * 0.07, H * 0.3 + 1.8, W * 0.48, 1.2)
                cv_motif = (W * 0.75, H * 0.52, W * 0.15, "w")
            else:
                deco = [(SH.ROUND_2_SAME_RECTANGLE, W - 1.7, 0, 1.0, 1.7, p["soft"], None, 1, 180, 0.5)]
                mo = (W - 1.2, 0.85, 0.8, None)
        elif st == 8:  # 우표
            bg = p["main"] if cover else p["soft"]
            if cover:
                deco = [(SH.RECTANGLE, 0.55, 0.55, W - 1.1, H - 1.1, "#FFFFFF", p["soft"], 3.0, 0, None, True)]
                cv_title = (W * 0.12, H * 0.33, W * 0.76, 1.5); cv_sub = (W * 0.12, H * 0.33 + 1.6, W * 0.76, 1.2)
                cv_motif = (W * 0.5, H * 0.21, H * 0.24, None)
            else:
                deco = [(SH.RECTANGLE, 0.3, 0.3, W - 0.6, H - 0.6, "#FFFFFF", p["main"], 2.0, 0, None, True)]
                tb = (0.85, 0.6, W - 3.1, 0.9); rect = (0.85, 1.7, W - 0.85, H - 0.7); mo = (W - 1.3, 1.05, 1.0, None)
        elif st == 9:  # 리본
            bg = "#FFFFFF"
            if cover:
                deco = [(SH.PENTAGON, 0, H * 0.3, W * 0.84, 2.0, p["main"]), (SH.RECTANGLE, 0, H * 0.3 + 2.0, W * 0.78, 0.07, p["accent"])]
                cv_title = (W * 0.06, H * 0.3 + 0.12, W * 0.72, 1.76); cv_sub = (W * 0.06, H * 0.3 + 2.3, W * 0.7, 1.2)
                cv_motif = (W * 0.88, H * 0.2, H * 0.34, None)
            else:
                deco = [(SH.PENTAGON, 0, 0.42, W * 0.78, 0.98, p["main"])]
                tb = (0.7, 0.45, W * 0.78 - 1.2, 0.9); rect = (0.7, 1.8, W - 0.7, H - 0.5); mo = (W - 0.95, 0.85, 1.1, None)
        if cover:
            tc = "#FFFFFF" if (st in (9,) or (st == 0 and False)) else tc
            if st == 9:
                tc = "#FFFFFF"
        else:
            if st == 9:
                tc = "#FFFFFF"
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg)
        for d in deco:
            kind, x, y, w, h = d[:5]
            fill = d[5] if len(d) > 5 else None
            line = d[6] if len(d) > 6 else None
            lw = d[7] if len(d) > 7 else 1.0
            rot = d[8] if len(d) > 8 else 0
            adj = d[9] if len(d) > 9 else None
            dash = d[10] if len(d) > 10 else False
            self.shape(kind, x, y, w, h, fill=fill, line=line, lw=lw, rot=rot, adj=adj, dash=dash, name="장식", s=s)
        self.title_color, self.rect = tc, rect
        if cover:
            self.tbox, self.sbox = cv_title, cv_sub
            mx, my, ms, flag = cv_motif
            cols = ("#FFFFFF", p["soft"], p["accent"]) if flag == "w" else None
            self.motif(mx, my, ms, cols)
        else:
            self.tbox = tb
            mx, my, ms, _ = mo
            if not self.cover_img:  # 그림 표지를 쓰는 덱은 안쪽 쪽에 클립아트를 넣지 않습니다
                self.motif(mx, my, ms)
