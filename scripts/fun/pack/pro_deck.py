"""PPT 한 벌(Deck) — 디자인 언어(pro_design)를 레이아웃으로 심고, 항목 모듈이 쓰는 쪽·블록(카드·줄·타임라인·사진·표·차트)을 제공합니다."""
import pack_kit as pk
import pro_design as D
from pro_kit import Pro
from pptx.enum.shapes import MSO_SHAPE as SH

TRANS = ["fade", "push:l", "wipe:r", "cover:l", "pull:r", "split:horz", "push:u", "wedge"]
ANIM = ["fade", "wipe-l", "wipe-u"]
GEOM_BY_PF = ["oval", "arch", "round"]
DEFAULT_NOTES = "글은 눌러서 바로 고치세요. 사진 자리는 가운데 아이콘을 누르면 사진을 넣을 수 있어요. 색·글꼴은 [디자인 > 변형]에서 한 번에 바꿀 수 있어요."


class Deck(Pro):
    def __init__(self, i, k, item_en, title, size="wide", subject="", keywords=""):
        self.cfg = D.design_of(i, k)
        pal = D.palette(i)
        super().__init__(pal, size, title, subject, keywords, fonts=self.cfg["font"])
        self.i, self.k = i, k
        self.pfill = "white" if self.cfg["tone"] == "tint" else "soft"
        self.eyebrow = f"{D.SEASON_EN[pal['season']]}  ·  {item_en}"
        self.vw = D.short(pal["name"])
        self.rect = (0.8, 2.0, self.W - 0.8, self.H - 0.75)
        self._anim, self._links, self._sec, self._prev = [], [], None, None
        self.begin_layouts()
        D.cover_layout(self)
        D.cover_layout(self, closing=True)
        D.body_layouts(self)
        self.end_layouts()

    # ── 쪽 만들기 ──────────────────────────────────────
    def _flush(self):
        if self._anim and self._prev is not None:
            self.animate(self._prev, self._anim, ANIM[self.cfg["anim"]])
        self._anim = []

    def _section(self, name):
        if name != self._sec:
            self._sec = name
            return name
        return None

    def cover(self, title, sub, extra=None, notes=None, closing=False):
        self._flush()
        notes = notes or DEFAULT_NOTES
        s = self.add("closing" if closing else "cover", title=title, label=self.eyebrow, sub=sub, extra=extra, notes=notes,
                     section=self._section("마무리" if closing else "시작"), vword=self.vw)
        self.drop_unused_pics()
        self.transition(s, "fade")
        self._prev = s
        return s

    def page(self, title, label, notes=None, photo=False, quote=False):
        self._flush()
        notes = notes or DEFAULT_NOTES
        key = "photo" if photo else ("quote" if quote else "body")
        s = self.add(key, title=title, label=label, notes=notes, section=self._section("본문"), vword=self.vw)
        self.transition(s, TRANS[self.cfg["trans"]])
        self._prev = s
        return s

    def done(self):
        self._flush()
        for shp, idx in self._links:
            tgt = shp.shapes[0] if hasattr(shp, "shapes") else shp  # 그룹은 첫 도형(카드 바탕)에 링크를 건다
            self.link_to(tgt, self.prs.slides[idx])

    def save(self, path):
        self.done()
        super().save(path)

    # ── 빌딩 블록 ──────────────────────────────────────
    def _box(self, x, y, w, h, name="카드"):
        cs = self.cfg["cs"]
        if cs == 0:
            return self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, fill="white", shadow="soft", adj={"adj": 6000}, name=name)
        if cs == 1:
            return self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, fill="tint", adj={"adj": 5000}, name=name)
        if cs == 2:
            return self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, fill=("alpha", "white", 55), line="main", lw=1.25, adj={"adj": 7000}, name=name)
        if cs == 3:
            b = self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, fill="white", shadow="thin", adj={"adj": 3000}, name=name)
            self.shape(SH.RECTANGLE, x, y + .12, .07, h - .24, fill="main", name="포인트 막대")
            return b
        b = self.shape(SH.ROUNDED_RECTANGLE, x, y, w, h, fill="white", shadow="thin", adj={"adj": 3000}, name=name)
        self.shape(SH.RECTANGLE, x + .02, y, w - .04, .09, fill="main", name="위쪽 띠")
        return b

    def badge(self, x, y, n, dia=0.5, size=15):
        bs = self.cfg["bs"]
        if bs == 2:
            self.text(x - .1, y - .05, dia + .5, dia + .1, f"{n:02d}", size=size + 8, bold=True, color="main", anchor="m", major=True, name=f"번호 {n}")
        elif bs == 1:
            self.shape(SH.OVAL, x, y, dia, dia, fill=None, line="main", lw=1.5, text=str(n), tsize=size, tbold=True, tcolor="deep", name=f"번호 {n}")
        elif bs == 3:
            self.shape(SH.ROUNDED_RECTANGLE, x, y, dia, dia, fill="main", adj={"adj": 25000}, text=str(n), tsize=size, tbold=True, tcolor="ink", name=f"번호 {n}")
        elif bs == 4:
            self.shape(SH.OVAL, x, y, dia, dia, fill="accent", text=str(n), tsize=size, tbold=True, tcolor="white", name=f"번호 {n}")
        else:
            self.shape(SH.OVAL, x, y, dia, dia, fill="main", text=str(n), tsize=size, tbold=True, tcolor="ink", name=f"번호 {n}")

    def cards(self, items, y=None, h=None, size=16, head_size=18, numbered=True):
        """[(제목, 내용)] 가로 카드 — 카드마다 그룹으로 묶어 선택 창·애니메이션이 한 덩어리로 움직입니다."""
        l, t, r, b = self.rect
        y = t if y is None else y
        h = min((b - t) if h is None else h, 3.9)
        n, gap = len(items), 0.3
        cw = (r - l - gap * (n - 1)) / n
        for k, (head, body) in enumerate(items):
            x = l + k * (cw + gap)
            with self.group(f"카드 {k + 1}") as g:
                self._box(x, y, cw, h, "카드 바탕")
                off = 0
                if numbered:
                    self.badge(x + .25, y + .28, k + 1)
                    off = .7
                self.text(x + .2, y + .25 + off, cw - .4, .7, head, size=head_size, bold=True, color="deep", anchor="m", major=True, name="카드 제목")
                self.text(x + .2, y + 1.0 + off, cw - .4, h - 1.15 - off, body, size=size, spacing=1.3, name="카드 내용")
            self._anim.append(g)

    def rows(self, items, y=None, h_each=0.95, size=15, head_size=18):
        l, t, r, b = self.rect
        y = t if y is None else y
        for k, (head, body) in enumerate(items):
            yy = y + k * (h_each + .15)
            with self.group(f"줄 {k + 1}") as g:
                self._box(l, yy, r - l, h_each, "줄 바탕")
                self.badge(l + .3, yy + (h_each - .5) / 2, k + 1)
                self.text(l + 1.1, yy + .05, 3.2, h_each - .1, head, size=head_size, bold=True, color="deep", anchor="m", major=True, name="줄 제목")
                self.text(l + 4.4, yy + .05, r - l - 4.7, h_each - .1, body, size=size, anchor="m", name="줄 내용")
            self._anim.append(g)

    def timeline(self, steps, y=None, size=15, head_size=18):
        l, t, r, b = self.rect
        y = (t + 1.6) if y is None else y
        n = len(steps)
        cw = (r - l) / n
        if self.cfg["tl"] == 0:
            self.line(l, y + .22, r, y + .22, "soft", 3, name="타임라인 선", tail="triangle")
        for k, (when, head, body) in enumerate(steps):
            cx = l + cw * k + cw / 2
            with self.group(f"때 {k + 1}") as g:
                if self.cfg["tl"] == 0:
                    self.shape(SH.OVAL, cx - .22, y, .44, .44, fill="main", line="white", lw=3, name="점")
                else:
                    self.shape(SH.PENTAGON, cx - cw / 2 + .05, y - .05, cw - .1, .54, fill="main" if k % 2 == 0 else "accent", name="단계 화살표", text=when, tsize=head_size - 3, tbold=True, tcolor="ink")
                if self.cfg["tl"] == 0:
                    self.text(cx - cw / 2 + .05, y - .7, cw - .1, .5, when, size=head_size - 2, bold=True, color="deep", align="c", anchor="b", major=True, name="때")
                self.text(cx - cw / 2 + .1, y + .65, cw - .2, .5, head, size=head_size, bold=True, color="deep", align="c", name="제목")
                self.text(cx - cw / 2 + .1, y + 1.2, cw - .2, 1.8, body, size=size, align="c", spacing=1.3, name="내용")
            self._anim.append(g)

    def two_col(self, left, right, y=None, h=None, size=15):
        l, t, r, b = self.rect
        y = t if y is None else y
        h = min(b - t, 3.9) if h is None else h
        gap = .4
        cw = (r - l - gap) / 2
        for k, (head, body) in enumerate((left, right)):
            x = l + k * (cw + gap)
            with self.group("왼쪽 칸" if k == 0 else "오른쪽 칸") as g:
                self._box(x, y, cw, h, "칸 바탕")
                self.shape(SH.ROUNDED_RECTANGLE, x + .2, y + .2, cw - .4, .65, fill="main" if k == 0 else "deep", adj={"adj": 50000}, name="칸 제목 알약",
                           text=head, tsize=17 if self.wide else 15, tbold=True, tcolor="ink" if k == 0 else "white")
                self.text(x + .3, y + 1.05, cw - .6, h - 1.2, body, size=size, bullet=True, spacing=1.3, space_after=6, name="칸 내용")
            self._anim.append(g)

    def person(self, x, y, w, h, role, lines, idx):
        """사진 자리 + 이름·소개 (사진 레이아웃의 그림 자리표시자 idx 를 씁니다)"""
        pd = min(w - .5, h * .5)
        geom = GEOM_BY_PF[self.cfg["pf"]]
        gw, gh = (pd, pd) if geom != "arch" else (pd * .86, pd)
        bh = min(h, gh + 2.3)
        self._box(x, y, w, bh, f"{role} 카드")
        self.photo_slot(idx, x + (w - gw) / 2, y + .25, gw, gh, geom=geom, label="사진", fill="white" if self.cfg["cs"] == 1 else "tint")
        self.text(x, y + gh + .35, w, .45, role, size=16, bold=True, color="deep", align="c", major=True, name=f"{role} 이름")
        self.text(x + .1, y + gh + .85, w - .2, bh - gh - .95, lines, size=12 if self.wide else 11, color="mute", align="c", spacing=1.25, name=f"{role} 소개")

    def photo_grid(self, n_cols, n_rows, labels=None, y=None, h=None, start_idx=13):
        l, t, r, b = self.rect
        y = t if y is None else y
        h = (b - t) if h is None else h
        gap = .22
        cw = (r - l - gap * (n_cols - 1)) / n_cols
        ch = (h - gap * (n_rows - 1)) / n_rows
        geom = ["round", "round", "arch", "oval"][self.cfg["pf"]] if n_cols * n_rows <= 6 else "round"
        k = 0
        for i in range(n_rows):
            for j in range(n_cols):
                g = geom if geom != "oval" else "round"
                lab = labels[k] if labels and k < len(labels) else f"사진 {k + 1}"
                self.photo_slot(start_idx + k, l + j * (cw + gap), y + i * (ch + gap), cw, ch, geom=g, label=lab)
                k += 1
        self.drop_unused_pics()

    def note(self, text, y=None, size=11):
        l, t, r, b = self.rect
        self.text(l, (b - .1) if y is None else y, r - l, .4, text, size=size, color="mute", name="안내")

    def statement(self, text):
        l, t, r, b = self.rect
        self.text(l, t, r - l, 1.3, text, size=26, bold=True, color="deep", align="c", anchor="m", major=True, name="핵심 문장")
